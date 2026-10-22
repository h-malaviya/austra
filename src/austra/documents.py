"""Protected document storage, text extraction, and search."""

import csv
import hashlib
import io
import re
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from austra.store import Store, _id, _now

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARS = 200_000
MAX_ARCHIVE_BYTES = 50 * 1024 * 1024
MAX_PDF_PAGES = 100
SUPPORTED_TYPES = {".pdf", ".docx", ".csv"}


def _extract_pdf(content: bytes) -> str:
    if not content.startswith(b"%PDF-"):
        raise ValueError("The file is not a PDF")
    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise ValueError("Encrypted PDFs are not supported")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError("PDF has too many pages")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(content: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if "word/document.xml" not in {entry.filename for entry in entries}:
                raise ValueError("The file is not a DOCX document")
            if len(entries) > 5000 or sum(entry.file_size for entry in entries) > MAX_ARCHIVE_BYTES:
                raise ValueError("DOCX expands beyond the allowed size")
    except zipfile.BadZipFile as error:
        raise ValueError("The file is not a DOCX document") from error
    try:
        document = Document(io.BytesIO(content))
    except PackageNotFoundError as error:
        raise ValueError("The file is not a DOCX document") from error
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _extract_csv(content: bytes) -> str:
    try:
        source = io.StringIO(content.decode("utf-8-sig"), newline="")
    except UnicodeDecodeError as error:
        raise ValueError("CSV must use UTF-8 encoding") from error
    rows = csv.reader(source)
    return "\n".join(" | ".join(cell for cell in row) for row in rows)


def extract_text(file_name: str, content: bytes) -> tuple[str, str]:
    suffix = Path(file_name).suffix.lower()
    if suffix not in SUPPORTED_TYPES:
        raise ValueError("Only PDF, DOCX, and CSV files are supported")
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("File must be between 1 byte and 10 MB")
    try:
        if suffix == ".pdf":
            extracted = _extract_pdf(content)
        elif suffix == ".docx":
            extracted = _extract_docx(content)
        else:
            extracted = _extract_csv(content)
    except (PdfReadError, csv.Error) as error:
        raise ValueError("File content could not be read") from error
    if len(extracted) > MAX_EXTRACTED_CHARS:
        raise ValueError("Extracted text is too large")
    status = "indexed" if extracted.strip() else "needs_text_review"
    return extracted, status


class DocumentLibrary:
    def __init__(self, store: Store, directory: str | Path) -> None:
        self.store = store
        self.directory = Path(directory)
        self.directory.mkdir(parents=True, exist_ok=True)

    def add(
        self, file_name: str, content: bytes, uploaded_by: str, case_id: str | None = None
    ) -> dict[str, Any]:
        if not file_name or len(file_name) > 255:
            raise ValueError("Invalid file name")
        text, status = extract_text(file_name, content)
        document_id = _id()
        suffix = Path(file_name).suffix.lower()
        stored_name = f"{document_id}{suffix}"
        location = self.directory / stored_name
        location.write_bytes(content)
        try:
            with self.store._connect() as connection:
                if case_id:
                    self.store._case_row(connection, case_id)
                connection.execute(
                    "INSERT INTO documents VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    (
                        document_id,
                        Path(file_name).name,
                        suffix,
                        hashlib.sha256(content).hexdigest(),
                        stored_name,
                        text,
                        status,
                        uploaded_by,
                        case_id,
                        _now(),
                    ),
                )
                connection.execute(
                    "INSERT INTO document_index VALUES (?, ?, ?)",
                    (document_id, Path(file_name).name, text),
                )
                self.store._event(
                    connection,
                    case_id,
                    "document_uploaded",
                    {
                        "document_id": document_id,
                        "file_name": Path(file_name).name,
                        "status": status,
                        "uploaded_by": uploaded_by,
                    },
                )
        except Exception:
            location.unlink(missing_ok=True)
            raise
        return {
            "id": document_id,
            "file_name": Path(file_name).name,
            "status": status,
            "case_id": case_id,
            "uploaded_by": uploaded_by,
            "size_bytes": len(content),
        }

    def list_documents(self) -> list[dict[str, Any]]:
        with self.store._connect() as connection:
            return [
                dict(row)
                for row in connection.execute(
                    "SELECT id, file_name, file_type, status, uploaded_by, case_id, created_at "
                    "FROM documents ORDER BY created_at DESC LIMIT 100"
                )
            ]

    def for_case(self, case_id: str) -> list[dict[str, Any]]:
        with self.store._connect() as connection:
            return [
                dict(row)
                for row in connection.execute(
                    "SELECT id, file_name, file_type, status, uploaded_by, case_id, created_at "
                    "FROM documents WHERE case_id = ? ORDER BY created_at",
                    (case_id,),
                )
            ]

    def get_text(self, document_id: str) -> dict[str, Any]:
        with self.store._connect() as connection:
            row = connection.execute(
                "SELECT id, file_name, status, extracted_text, uploaded_by, case_id, created_at "
                "FROM documents WHERE id = ?",
                (document_id,),
            ).fetchone()
            if row is None:
                raise KeyError(document_id)
            return dict(row)

    def search(self, query: str, limit: int = 5) -> list[dict[str, Any]]:
        terms = re.findall(r"[A-Za-z0-9]{3,}", query.lower())[:8]
        if not terms:
            return []
        expression = " OR ".join(terms)
        with self.store._connect() as connection:
            rows = connection.execute(
                "SELECT d.id, d.file_name, d.status, d.case_id, "
                "snippet(document_index, 2, '', '', ' ... ', 40) AS excerpt "
                "FROM document_index JOIN documents d ON d.id = document_index.document_id "
                "WHERE document_index MATCH ? ORDER BY bm25(document_index) LIMIT ?",
                (expression, limit),
            )
            return [dict(row) for row in rows]
