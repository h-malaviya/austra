"""Local chat interface for an initially blank business workspace."""

import asyncio
import hashlib
import io
import json
import logging
import os
import re
import secrets
import threading
import time
from collections.abc import Callable
from contextlib import asynccontextmanager, suppress
from datetime import UTC, datetime
from pathlib import Path
from typing import Annotated, Any, Literal

import httpx
from docx import Document
from fastapi import (
    BackgroundTasks,
    Depends,
    FastAPI,
    File,
    Form,
    HTTPException,
    Request,
    Response,
    UploadFile,
)
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from starlette.concurrency import run_in_threadpool

from austra.access import CONNECTORS, GUIDES, VISIBLE_FIELDS, Access
from austra.admin_intents import is_rule_suggestion, rule_listing, verification_update
from austra.agent import (
    AgentReply,
    AnalysisUnavailable,
    ConnectionCheckError,
    Draft,
    analyze,
    verify_anthropic_key,
)
from austra.connectors import VERIFIERS, Verifier, exchange_gmail_code, gmail_authorization_url
from austra.files import SECRET_PATTERN, Files
from austra.imports import import_chat_tables
from austra.infrastructure import Infrastructure, InfrastructureError
from austra.pdf_text import PDFExtractionError, extract_pdf_text
from austra.reports import (
    is_stock_request,
    requested_stock_unit,
    stock_display_unit,
    stock_report,
)
from austra.skills import PLATFORM_SKILLS
from austra.slack import event_message, verify_signature
from austra.slack_socket import run_socket_mode

logger = logging.getLogger(__name__)
STATIC_DIR = Path(__file__).parent / "static"
MAX_ATTACHMENT_BYTES = 5 * 1024 * 1024
MAX_ATTACHMENTS = 5
MAX_ATTACHMENT_TEXT_CHARS = 20000
DIRECT_TABLE_IMPORT_BYTES = 20 * 1024
SUPPORTED_ATTACHMENTS = {".pdf", ".docx", ".csv", ".txt", ".md", ".xlsx"}
SESSION_COOKIE = "austra_session"
OAUTH_STATE_SECONDS = 600


class Message(BaseModel):
    body: str = Field(min_length=1)

    @property
    def safe_body(self) -> str:
        if not self.body.strip():
            raise ValueError("Message cannot be blank")
        if SECRET_PATTERN.search(self.body):
            raise ValueError("Use the private Connections prompt for credentials")
        return self.body.strip()


class Passphrase(BaseModel):
    password: str = Field(min_length=1, max_length=200)


class Credentials(BaseModel):
    values: dict[str, str]


class Review(BaseModel):
    decision: Literal["approved", "rejected"]


def _extract_attachment(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_text(content)
    if suffix == ".docx":
        return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs)
    if suffix in {".csv", ".txt", ".md"}:
        return content.decode("utf-8-sig")
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            return "\n".join(
                "\t".join("" if value is None else str(value) for value in row)
                for sheet in workbook.worksheets
                for row in sheet.values
            )
        finally:
            workbook.close()
    raise ValueError("Only PDF, DOCX, CSV, TXT, MD, and XLSX files are supported")


async def _prepare_uploads(uploads: list[UploadFile]) -> list[dict[str, Any]]:
    if len(uploads) > MAX_ATTACHMENTS:
        raise HTTPException(status_code=400, detail=f"Attach up to {MAX_ATTACHMENTS} files at once")
    prepared = []
    for upload in uploads:
        filename = (upload.filename or "").replace("\\", "/").rsplit("/", 1)[-1]
        filename = re.sub(r"[\x00-\x1f\x7f]", " ", filename).strip()[:160]
        extension = Path(filename).suffix.lower()
        if not filename or extension not in SUPPORTED_ATTACHMENTS:
            raise HTTPException(
                status_code=400, detail="Only PDF, DOCX, CSV, TXT, MD, and XLSX files are supported"
            )
        content = await upload.read(MAX_ATTACHMENT_BYTES + 1)
        await upload.close()
        if len(content) > MAX_ATTACHMENT_BYTES:
            raise HTTPException(status_code=413, detail="Each attachment must be 5 MB or smaller")
        try:
            extracted = await run_in_threadpool(_extract_attachment, filename, content)
        except PDFExtractionError as error:
            raise HTTPException(status_code=400, detail=f"{filename}: {error}") from error
        except Exception as error:
            logger.warning(
                "Attachment extraction failed for %s: %s", filename, type(error).__name__
            )
            raise HTTPException(status_code=400, detail=f"Could not read {filename}") from error
        if not extracted.strip():
            detail = (
                f"No readable text was found in {filename}, even after scanning its pages. "
                "Try a clearer PDF."
                if extension == ".pdf"
                else f"{filename} has no readable text"
            )
            raise HTTPException(status_code=400, detail=detail)
        if SECRET_PATTERN.search(extracted):
            raise HTTPException(
                status_code=400, detail="Attachment appears to contain a credential"
            )
        prepared.append(
            {
                "id": hashlib.sha256(content).hexdigest(),
                "name": filename,
                "extension": extension,
                "content": content,
                "text": extracted,
                "excerpt": ""
                if extension in {".csv", ".xlsx"} and len(extracted) > MAX_ATTACHMENT_TEXT_CHARS
                else extracted[:MAX_ATTACHMENT_TEXT_CHARS],
            }
        )
    return prepared


def create_app(
    root: str | Path | None = None,
    planner: Callable[[Files, Access, str], AgentReply] = analyze,
    anthro_verifier: Callable[[str], None] = verify_anthropic_key,
    connector_verifiers: dict[str, Verifier] | None = None,
) -> FastAPI:
    workspace = Path(root or os.getenv("AUSTRA_WORKSPACE", "data/blank"))
    files = Files(workspace)
    access = Access(workspace)
    infrastructure = Infrastructure()
    verifiers = {**VERIFIERS, **(connector_verifiers or {})}
    oauth_states: dict[str, tuple[float, str]] = {}
    jobs: dict[str, dict[str, Any]] = {}
    jobs_lock = threading.RLock()
    slack_socket_status = {"connected": False}
    public_url = os.getenv("AUSTRA_PUBLIC_URL", "http://127.0.0.1:8000").rstrip("/")
    gmail_redirect_uri = f"{public_url}/api/connections/gmail/callback"

    @asynccontextmanager
    async def lifespan(_app: FastAPI):
        task = asyncio.create_task(
            run_socket_mode(access, files, process_socket_event, slack_socket_status)
        )
        try:
            yield
        finally:
            task.cancel()
            with suppress(asyncio.CancelledError):
                await task

    app = FastAPI(title="Austra", lifespan=lifespan)
    app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

    @app.middleware("http")
    async def restrict_local_access(request: Request, call_next):
        if (
            request.url.path != "/integrations/slack/events"
            and os.getenv("AUSTRA_DOCKER_PROXY") != "1"
        ):
            host = request.client.host if request.client else ""
            if host not in {"127.0.0.1", "::1", "testclient"}:
                return JSONResponse({"detail": "Local access only"}, status_code=403)
        response = await call_next(request)
        if request.url.path == "/" or request.url.path.startswith(("/static/", "/api/")):
            response.headers["Cache-Control"] = "no-store"
        return response

    def require_admin(request: Request) -> None:
        if not access.authenticated(request.cookies.get(SESSION_COOKIE)):
            raise HTTPException(status_code=401, detail="Administrator sign in required")

    def analyzed_chat(
        chat_id: str, progress: Callable[[str], None] | None = None
    ) -> dict[str, Any]:
        latest = files.chat(chat_id)["messages"][-1]
        latest_text = latest.get("display_body", latest["body"])
        table_attachments = [
            attachment
            for message in files.chat(chat_id)["messages"]
            if message["role"] == "user"
            for attachment in message.get("attachments", [])
            if attachment["name"].lower().endswith((".csv", ".xlsx"))
        ]
        import_intent = bool(
            re.search(r"\b(add|import|save|ingest|load|continue)\b", latest_text, re.IGNORECASE)
        )
        large_table = any(
            attachment["size"] > DIRECT_TABLE_IMPORT_BYTES for attachment in table_attachments
        )
        if (
            table_attachments
            and (large_table or planner is analyze)
            and (import_intent or not latest_text.strip())
        ):
            if progress:
                progress("Inspecting table structures")
            imported, issues = import_chat_tables(
                files, chat_id, auto_schema=True, progress=progress
            )
            reply = (
                f"Saved {imported} business record(s)."
                if imported
                else "No new rows needed saving."
            )
            if issues:
                reply += "\n\n" + "\n".join(issues)
            files.add_message(chat_id, "assistant", reply, "Austra")
            files.record_event("agent.replied", "Austra", "chat", chat_id, chat_id=chat_id)
            return {
                "chat": files.chat(chat_id),
                "drafts": files.drafts(chat_id),
                "credential_requests": [],
                "analysis_error": "\n".join(issues) or None,
            }
        rules = rule_listing(files, chat_id)
        if rules is not None:
            files.add_message(chat_id, "assistant", rules, "Austra")
            files.record_event("agent.replied", "Austra", "chat", chat_id, chat_id=chat_id)
            return {
                "chat": files.chat(chat_id),
                "drafts": files.drafts(chat_id),
                "credential_requests": [],
                "analysis_error": None,
            }
        unit_request = requested_stock_unit(latest_text)
        if latest["role"] == "user" and unit_request and files.chat(chat_id)["source"] == "ui":
            label = "metric tons (t)" if unit_request == "t" else "kilograms (kg)"
            files.save_policy(
                chat_id,
                "stock_display_units",
                "# Stock display units\n\n"
                f"Display unit: {label}\n\n"
                "Display stock and inventory mass in this unit. Store underlying quantities in kilograms.",
                "Administrator set the stock display unit",
            )
            files.add_message(
                chat_id,
                "assistant",
                f"Stock and inventory quantities will display in {label}. Saved values remain in kilograms.",
                "Austra",
            )
            files.record_event("agent.replied", "Austra", "chat", chat_id, chat_id=chat_id)
            return {
                "chat": files.chat(chat_id),
                "drafts": files.drafts(chat_id),
                "credential_requests": [],
                "analysis_error": None,
            }
        if (
            latest["role"] == "user"
            and not latest.get("attachments")
            and is_stock_request(latest_text)
        ):
            files.add_message(
                chat_id,
                "assistant",
                stock_report(files.records(), stock_display_unit(files.approved())),
                "Austra",
            )
            files.record_event("agent.replied", "Austra", "chat", chat_id, chat_id=chat_id)
            return {
                "chat": files.chat(chat_id),
                "drafts": files.drafts(chat_id),
                "credential_requests": [],
                "analysis_error": None,
            }
        analysis_error = None
        try:
            if progress:
                progress("Analyzing request")
            result = verification_update(files, chat_id) or planner(files, access, chat_id)
        except AnalysisUnavailable as error:
            logger.warning("Chat analysis unavailable: %s", error)
            files.record_event(
                "agent.analysis_failed",
                "Austra",
                "chat",
                chat_id,
                chat_id=chat_id,
                skills_used=list(PLATFORM_SKILLS),
                outcome="failed",
            )
            analysis_error = str(error)
            result = AgentReply(reply=str(error))
        if result.title.strip() and not files.chat(chat_id).get("title_generated"):
            files.set_chat_title(chat_id, result.title)
        if is_rule_suggestion(latest_text) and result.rule_changes:
            result.drafts.extend(
                Draft(
                    kind="policy",
                    name=change.name,
                    summary=change.reason,
                    content=change.content,
                )
                for change in result.rule_changes
                if change.action == "activate" and change.content.strip()
            )
            result.rule_changes = []
            result.reply = "I drafted the suggested rule for review. It is not active."
        if result.decision == "deny":
            files.record_event(
                "agent.request_denied",
                "Austra",
                "chat",
                chat_id,
                chat_id=chat_id,
                skills_used=list(PLATFORM_SKILLS),
            )
            result = AgentReply(reply=result.reply, decision="deny")
        elif result.drafts or result.save_records or result.record_changes or result.rule_changes:
            try:
                if (result.record_changes or result.rule_changes) and files.chat(chat_id)[
                    "source"
                ] != "ui":
                    raise ValueError("Changes require an administrator chat in the app")
                has_table = any(
                    attachment["name"].lower().endswith((".csv", ".xlsx"))
                    for message in files.chat(chat_id)["messages"]
                    for attachment in message.get("attachments", [])
                )
                applied = files.apply_agent_drafts(
                    chat_id,
                    [
                        draft.model_dump()
                        for draft in result.drafts
                        if not (has_table and result.save_records and draft.kind == "record")
                    ],
                    list(PLATFORM_SKILLS),
                )
                imported, issues = (
                    import_chat_tables(files, chat_id, auto_schema=True, progress=progress)
                    if result.save_records
                    else (0, [])
                )
                saved_records = sum(item["kind"] == "record" for item in applied) + imported
                updates = []
                removed_records = 0
                updated_records = 0
                for change in result.record_changes:
                    if change.action == "delete":
                        identifier = (
                            change.record_id if change.scope == "id" else change.schema_name
                        )
                        removed_records += files.delete_records(
                            chat_id, change.scope, identifier, change.reason
                        )
                    else:
                        values = json.loads(change.values_json)
                        if not isinstance(values, dict):
                            raise ValueError("Updated record must be a JSON object")
                        if change.scope == "id":
                            files.update_record(chat_id, change.record_id, values, change.reason)
                            updated_records += 1
                        elif change.scope == "schema":
                            updated_records += files.update_records(
                                chat_id,
                                change.schema_name,
                                values,
                                change.reason,
                                change.source_attachment,
                            )
                        else:
                            raise ValueError("Bulk updates require a named record schema")
                activated_rules = 0
                retired_rules = 0
                for change in result.rule_changes:
                    if change.action == "activate":
                        files.save_policy(chat_id, change.name, change.content, change.reason)
                        activated_rules += 1
                    else:
                        files.retire_policy(chat_id, change.name, change.reason)
                        retired_rules += 1
                if any(item["kind"] == "schema" for item in applied):
                    updates.append("The business schema is ready to use.")
                if saved_records:
                    updates.append(f"Saved {saved_records} business record(s).")
                elif result.save_records and not issues and has_table:
                    updates.append("No new rows needed saving.")
                if updated_records:
                    updates.append(f"Updated {updated_records} business record(s).")
                if activated_rules:
                    updates.append(f"Activated {activated_rules} business rule(s).")
                if retired_rules:
                    updates.append(f"Retired {retired_rules} business rule(s).")
                if removed_records:
                    updates.append(
                        f"Removed {removed_records} business record(s) from active data. "
                        "Audit copies are retained."
                    )
                elif any(change.action == "delete" for change in result.record_changes):
                    updates.append("No matching active records remained to remove.")
                updates.extend(issues)
                if updates:
                    result.reply = f"{result.reply}\n\n{' '.join(updates)}"
            except ValueError as error:
                files.record_event(
                    "agent.draft_rejected",
                    "Austra",
                    "chat",
                    chat_id,
                    chat_id=chat_id,
                    skills_used=list(PLATFORM_SKILLS),
                    outcome="failed",
                )
                analysis_error = str(error)
                result = AgentReply(reply=f"I could not save the data: {error}")
        files.add_message(chat_id, "assistant", result.reply, "Austra")
        files.record_event(
            "agent.replied",
            "Austra",
            "chat",
            chat_id,
            chat_id=chat_id,
            skills_used=list(PLATFORM_SKILLS),
        )
        return {
            "chat": files.chat(chat_id),
            "drafts": files.drafts(chat_id),
            "credential_requests": result.credential_requests,
            "analysis_error": analysis_error,
        }

    def submitted_chat(
        chat_id: str | None, body: str, uploads: list[dict[str, Any]]
    ) -> dict[str, Any]:
        chat, duplicates, changed = submit_only(chat_id, body, uploads)
        result = (
            analyzed_chat(chat["id"])
            if changed
            else {
                "chat": chat,
                "drafts": files.drafts(chat["id"]),
                "credential_requests": [],
                "analysis_error": None,
            }
        )
        result["duplicate_files"] = duplicates
        result["unchanged"] = not changed
        return result

    def submit_only(
        chat_id: str | None, body: str, uploads: list[dict[str, Any]]
    ) -> tuple[dict[str, Any], list[str], bool]:
        try:
            text = Message(body=body).safe_body if body.strip() else ""
            return files.submit_message(chat_id, text, uploads)
        except KeyError as error:
            raise HTTPException(status_code=404, detail="Chat not found") from error
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error

    def start_job(chat_id: str) -> dict[str, Any]:
        with jobs_lock:
            if any(
                job["chat_id"] == chat_id and job["status"] == "running" for job in jobs.values()
            ):
                raise HTTPException(
                    status_code=409, detail="A process is already running for this chat"
                )
            job_id = secrets.token_hex(16)
            job = {
                "id": job_id,
                "chat_id": chat_id,
                "chat_title": files.chat(chat_id)["title"],
                "status": "running",
                "stage": "Starting",
                "started_at": datetime.now(UTC).isoformat(),
                "finished_at": None,
                "error": None,
            }
            jobs[job_id] = job

        def progress(stage: str) -> None:
            with jobs_lock:
                job["stage"] = stage

        def run() -> None:
            try:
                result = analyzed_chat(chat_id, progress)
                with jobs_lock:
                    job["status"] = "failed" if result["analysis_error"] else "completed"
                    job["error"] = result["analysis_error"]
                    job["stage"] = "Finished" if not result["analysis_error"] else "Needs attention"
            except Exception:
                logger.exception("Background chat processing failed")
                with jobs_lock:
                    job["status"] = "failed"
                    job["error"] = "The process failed. Retry this chat."
                    job["stage"] = "Needs attention"
            finally:
                with jobs_lock:
                    job["finished_at"] = datetime.now(UTC).isoformat()

        threading.Thread(target=run, daemon=True, name=f"austra-job-{job_id}").start()
        return job_snapshot(job)

    def job_snapshot(job: dict[str, Any]) -> dict[str, Any]:
        started = datetime.fromisoformat(job["started_at"])
        finished = (
            datetime.fromisoformat(job["finished_at"]) if job["finished_at"] else datetime.now(UTC)
        )
        return {**job, "elapsed_seconds": max(0, int((finished - started).total_seconds()))}

    @app.get("/")
    def home() -> FileResponse:
        return FileResponse(STATIC_DIR / "index.html")

    @app.get("/api/config")
    def config(request: Request) -> dict[str, Any]:
        return {
            "admin_ready": access.ready,
            "authenticated": access.authenticated(request.cookies.get(SESSION_COOKIE)),
            "connectors": access.configured(),
            "guides": GUIDES,
            "missing": {name: access.missing(name) for name in CONNECTORS},
            "fields": {
                name: list(VISIBLE_FIELDS.get(name, fields)) for name, fields in CONNECTORS.items()
            },
            "gmail_redirect_uri": gmail_redirect_uri,
            "slack_socket_connected": slack_socket_status["connected"],
        }

    @app.post("/api/setup")
    def setup(payload: Passphrase, response: Response) -> dict[str, bool]:
        try:
            session = access.setup(payload.password)
        except ValueError as error:
            files.record_event(
                "admin.setup_failed", "unknown", "access", "administrator", outcome="failed"
            )
            raise HTTPException(status_code=400, detail=str(error)) from error
        response.set_cookie(SESSION_COOKIE, session, httponly=True, samesite="strict")
        files.record_event("admin.setup", "Admin", "access", "administrator")
        return {"authenticated": True}

    @app.post("/api/login")
    def login(payload: Passphrase, response: Response) -> dict[str, bool]:
        try:
            session = access.login(payload.password)
        except ValueError as error:
            files.record_event(
                "admin.login_failed", "unknown", "access", "administrator", outcome="failed"
            )
            raise HTTPException(status_code=401, detail=str(error)) from error
        response.set_cookie(SESSION_COOKIE, session, httponly=True, samesite="strict")
        files.record_event("admin.login", "Admin", "access", "administrator")
        return {"authenticated": True}

    @app.post("/api/logout")
    def logout(request: Request, response: Response) -> dict[str, bool]:
        access.logout(request.cookies.get(SESSION_COOKIE))
        response.delete_cookie(SESSION_COOKIE)
        files.record_event("admin.logout", "Admin", "access", "administrator")
        return {"authenticated": False}

    @app.post("/api/credentials/{connector}", dependencies=[Depends(require_admin)])
    def save_credentials(connector: str, payload: Credentials) -> dict[str, Any]:
        try:
            if connector == "anthropic":
                key = payload.values.get("ANTHROPIC_API_KEY", "")
                if not key:
                    raise ValueError("Enter an Anthropic API key")
                if len(key) > 8000 or "\n" in key or "\r" in key:
                    raise ValueError("Credential values must be nonempty single lines")
                anthro_verifier(key)
            missing = access.save(connector, payload.values)
            files.record_event("credential.updated", "Admin", "connector", connector)
        except ConnectionCheckError as error:
            raise HTTPException(status_code=error.status_code, detail=str(error)) from error
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error
        return {"connector": connector, "configured": not missing, "missing": missing}

    @app.post("/api/credentials/anthropic/test", dependencies=[Depends(require_admin)])
    def test_anthropic_connection() -> dict[str, bool]:
        return test_connection("anthropic")

    @app.post("/api/connections/{connector}/test", dependencies=[Depends(require_admin)])
    def test_connection(connector: str) -> dict[str, bool]:
        if connector not in CONNECTORS:
            raise HTTPException(status_code=404, detail="Unknown connection")
        missing = access.missing(connector)
        if missing:
            raise HTTPException(status_code=400, detail="Complete all connection fields first")
        try:
            if connector == "anthropic":
                anthro_verifier(access.get("ANTHROPIC_API_KEY"))
            else:
                verifiers[connector]({name: access.get(name) for name in CONNECTORS[connector]})
        except ConnectionCheckError as error:
            files.record_event(
                "connection.tested", "Admin", "connector", connector, outcome="failed"
            )
            raise HTTPException(status_code=error.status_code, detail=str(error)) from error
        files.record_event("connection.tested", "Admin", "connector", connector)
        return {"connected": True}

    @app.post("/api/connections/gmail/start", dependencies=[Depends(require_admin)])
    def start_gmail_connection(request: Request) -> dict[str, str]:
        client_id = access.get("GMAIL_CLIENT_ID")
        if not client_id or not access.get("GMAIL_CLIENT_SECRET"):
            raise HTTPException(status_code=400, detail="Save the Gmail OAuth client details first")
        session = request.cookies.get(SESSION_COOKIE, "")
        state = secrets.token_urlsafe(32)
        oauth_states[state] = (time.monotonic() + OAUTH_STATE_SECONDS, session)
        files.record_event("connection.authorization_started", "Admin", "connector", "gmail")
        return {"authorization_url": gmail_authorization_url(client_id, gmail_redirect_uri, state)}

    @app.get("/api/connections/gmail/callback")
    def finish_gmail_connection(request: Request) -> RedirectResponse:
        state = request.query_params.get("state", "")
        pending = oauth_states.pop(state, None)
        if not pending or pending[0] < time.monotonic() or not access.authenticated(pending[1]):
            files.record_event(
                "connection.authorization_failed", "unknown", "connector", "gmail", outcome="failed"
            )
            raise HTTPException(status_code=400, detail="Gmail authorization expired. Try again")
        result = "denied"
        code = request.query_params.get("code", "")
        if code and len(code) <= 2048 and not request.query_params.get("error"):
            try:
                refresh_token, email = exchange_gmail_code(
                    code,
                    access.get("GMAIL_CLIENT_ID"),
                    access.get("GMAIL_CLIENT_SECRET"),
                    gmail_redirect_uri,
                )
                access.save(
                    "gmail", {"GMAIL_REFRESH_TOKEN": refresh_token, "GMAIL_SENDER_EMAIL": email}
                )
                files.record_event("credential.updated", "Admin", "connector", "gmail")
                files.record_event("connection.tested", "Admin", "connector", "gmail")
                result = "connected"
            except (ConnectionCheckError, ValueError):
                files.record_event(
                    "connection.authorization_failed",
                    "Admin",
                    "connector",
                    "gmail",
                    outcome="failed",
                )
                result = "failed"
        else:
            files.record_event(
                "connection.authorization_failed", "Admin", "connector", "gmail", outcome="failed"
            )
        response = RedirectResponse(f"/?connection=gmail&oauth={result}", status_code=303)
        response.set_cookie(SESSION_COOKIE, pending[1], httponly=True, samesite="strict")
        response.headers["Cache-Control"] = "no-store"
        response.headers["Referrer-Policy"] = "no-referrer"
        return response

    @app.get("/api/chats", dependencies=[Depends(require_admin)])
    def list_chats() -> list[dict[str, Any]]:
        return files.chats()

    @app.post("/api/chats", dependencies=[Depends(require_admin)], status_code=201)
    def create_chat(payload: Message) -> dict[str, Any]:
        try:
            body = payload.safe_body
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error
        chat = files.create_chat(body, "Admin")
        return analyzed_chat(chat["id"])

    @app.get("/api/chats/{chat_id}", dependencies=[Depends(require_admin)])
    def get_chat(chat_id: str) -> dict[str, Any]:
        try:
            return {"chat": files.chat(chat_id), "drafts": files.drafts(chat_id)}
        except KeyError as error:
            raise HTTPException(status_code=404, detail="Chat not found") from error

    @app.get(
        "/api/chats/{chat_id}/attachments/{attachment_id}", dependencies=[Depends(require_admin)]
    )
    def get_attachment(chat_id: str, attachment_id: str) -> FileResponse:
        try:
            path, name = files.attachment(chat_id, attachment_id)
        except KeyError as error:
            raise HTTPException(status_code=404, detail="Attachment not found") from error
        return FileResponse(path, filename=name)

    @app.delete("/api/chats/{chat_id}", dependencies=[Depends(require_admin)])
    def delete_chat(chat_id: str) -> dict[str, bool]:
        try:
            files.delete_chat(chat_id)
        except KeyError as error:
            raise HTTPException(status_code=404, detail="Chat not found") from error
        return {"deleted": True, "audit_preserved": True}

    @app.post("/api/chats/{chat_id}/messages", dependencies=[Depends(require_admin)])
    def add_message(chat_id: str, payload: Message) -> dict[str, Any]:
        try:
            files.add_message(chat_id, "user", payload.safe_body, "Admin")
        except KeyError as error:
            raise HTTPException(status_code=404, detail="Chat not found") from error
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error
        return analyzed_chat(chat_id)

    @app.post("/api/chats/{chat_id}/retry", dependencies=[Depends(require_admin)])
    def retry_chat(chat_id: str) -> dict[str, Any]:
        try:
            files.chat(chat_id)
        except KeyError as error:
            raise HTTPException(status_code=404, detail="Chat not found") from error
        return analyzed_chat(chat_id)

    @app.post("/api/messages", dependencies=[Depends(require_admin)])
    async def submit_message(
        body: str = Form(default=""),
        chat_id: str = Form(default=""),
        uploads: Annotated[list[UploadFile] | None, File()] = None,
    ) -> dict[str, Any]:
        prepared = await _prepare_uploads(uploads or [])
        return submitted_chat(chat_id or None, body, prepared)

    @app.post("/api/messages/background", dependencies=[Depends(require_admin)], status_code=202)
    async def submit_background_message(
        body: str = Form(default=""),
        chat_id: str = Form(default=""),
        uploads: Annotated[list[UploadFile] | None, File()] = None,
    ) -> dict[str, Any]:
        if chat_id:
            with jobs_lock:
                if any(
                    job["chat_id"] == chat_id and job["status"] == "running"
                    for job in jobs.values()
                ):
                    raise HTTPException(
                        status_code=409, detail="A process is already running for this chat"
                    )
        prepared = await _prepare_uploads(uploads or [])
        chat, duplicates, changed = submit_only(chat_id or None, body, prepared)
        return {
            "chat": chat,
            "duplicate_files": duplicates,
            "unchanged": not changed,
            "job": start_job(chat["id"]) if changed else None,
        }

    @app.get("/api/jobs", dependencies=[Depends(require_admin)])
    def list_jobs() -> list[dict[str, Any]]:
        with jobs_lock:
            return [job_snapshot(job) for job in jobs.values() if job["status"] == "running"]

    @app.get("/api/jobs/{job_id}", dependencies=[Depends(require_admin)])
    def get_job(job_id: str) -> dict[str, Any]:
        with jobs_lock:
            if job_id not in jobs:
                raise HTTPException(status_code=404, detail="Process not found")
            return job_snapshot(jobs[job_id])

    @app.post("/api/chats/{chat_id}/attachments", dependencies=[Depends(require_admin)])
    async def add_attachment(chat_id: str, upload: Annotated[UploadFile, File()]) -> dict[str, Any]:
        prepared = await _prepare_uploads([upload])
        return submitted_chat(chat_id, "", prepared)

    @app.get("/api/artifacts", dependencies=[Depends(require_admin)])
    def list_artifacts() -> list[dict[str, Any]]:
        return files.approved()

    @app.get("/api/rules", dependencies=[Depends(require_admin)])
    def list_rules(query: str = "") -> list[dict[str, Any]]:
        match = query.casefold()
        return [
            item
            for item in files.approved()
            if item["kind"] in {"policy", "workflow"}
            and (
                not match or match in item["name"].casefold() or match in item["content"].casefold()
            )
        ]

    @app.get("/api/records", dependencies=[Depends(require_admin)])
    def list_records() -> list[dict[str, Any]]:
        return files.records()

    @app.get("/api/audit/records", dependencies=[Depends(require_admin)])
    def audit_records() -> list[dict[str, Any]]:
        return files.records(include_deleted=True)

    @app.get("/api/audit", dependencies=[Depends(require_admin)])
    def list_audit() -> list[dict[str, Any]]:
        return files.audit_entries()

    @app.get("/api/audit/chats/{chat_id}", dependencies=[Depends(require_admin)])
    def audit_chat(chat_id: str) -> dict[str, Any]:
        try:
            return files.audit_chat(chat_id)
        except KeyError as error:
            raise HTTPException(status_code=404, detail="Chat not found") from error

    @app.get(
        "/api/audit/chats/{chat_id}/attachments/{attachment_id}",
        dependencies=[Depends(require_admin)],
    )
    def audit_attachment(chat_id: str, attachment_id: str) -> FileResponse:
        try:
            path, name = files.attachment(chat_id, attachment_id, audit=True)
        except KeyError as error:
            raise HTTPException(status_code=404, detail="Attachment not found") from error
        return FileResponse(path, filename=name)

    @app.get("/api/infrastructure", dependencies=[Depends(require_admin)])
    def infrastructure_status() -> dict[str, Any]:
        try:
            return {"services": infrastructure.status()}
        except InfrastructureError as error:
            raise HTTPException(status_code=503, detail=str(error)) from error

    @app.post("/api/drafts/{draft_id}/review", dependencies=[Depends(require_admin)])
    def review_draft(draft_id: str, payload: Review) -> dict[str, Any]:
        try:
            return files.review(draft_id, payload.decision, infrastructure)
        except KeyError as error:
            raise HTTPException(status_code=404, detail="Draft not found") from error
        except ValueError as error:
            raise HTTPException(status_code=409, detail=str(error)) from error
        except InfrastructureError as error:
            raise HTTPException(status_code=503, detail=str(error)) from error

    def analyze_slack(chat_id: str, channel: str, thread_ts: str) -> None:
        result = analyzed_chat(chat_id)
        token = access.get("SLACK_BOT_TOKEN")
        if not token:
            return
        try:
            response = httpx.post(
                "https://slack.com/api/chat.postMessage",
                headers={"Authorization": f"Bearer {token}"},
                json={
                    "channel": channel,
                    "thread_ts": thread_ts,
                    "text": result["chat"]["messages"][-1]["body"],
                },
                timeout=15,
            )
            response.raise_for_status()
            if not response.json().get("ok"):
                logger.warning("Slack did not accept the reply")
                files.record_event(
                    "connection.reply_failed",
                    "Austra",
                    "connector",
                    "slack",
                    chat_id=chat_id,
                    outcome="failed",
                )
            else:
                files.record_event(
                    "connection.reply_sent", "Austra", "connector", "slack", chat_id=chat_id
                )
        except (httpx.HTTPError, ValueError):
            logger.exception("Slack reply failed")
            files.record_event(
                "connection.reply_failed",
                "Austra",
                "connector",
                "slack",
                chat_id=chat_id,
                outcome="failed",
            )

    def ingest_slack_event(payload: dict[str, Any]) -> tuple[str, str, str] | None:
        allowed_team = access.get("SLACK_ALLOWED_TEAM_ID")
        if allowed_team and payload.get("team_id") != allowed_team:
            files.record_event(
                "connection.event_rejected",
                "unknown",
                "connector",
                "slack",
                details={"reason": "workspace"},
                outcome="failed",
            )
            raise HTTPException(status_code=403, detail="Slack workspace not allowed")
        message = event_message(payload)
        if message is None:
            return None
        body_text, sender, event_id = message
        if SECRET_PATTERN.search(body_text):
            return None
        event = payload["event"]
        channel = event.get("channel", "")
        thread_ts = event.get("thread_ts") or event.get("ts", "")
        if (
            not isinstance(channel, str)
            or not isinstance(thread_ts, str)
            or not channel
            or not thread_ts
        ):
            raise HTTPException(status_code=400, detail="Invalid Slack thread")
        external_thread = f"{channel}:{thread_ts}"
        chat, fresh = files.ingest_external(event_id, external_thread, body_text, sender)
        if not fresh:
            return None
        files.record_event(
            "connection.event_received", sender, "connector", "slack", chat_id=chat["id"]
        )
        return chat["id"], channel, thread_ts

    def process_socket_event(payload: dict[str, Any]) -> None:
        try:
            delivery = ingest_slack_event(payload)
        except HTTPException:
            return
        if delivery:
            analyze_slack(*delivery)

    @app.post("/integrations/slack/events")
    async def slack_events(request: Request, background: BackgroundTasks) -> JSONResponse:
        signing_secret = access.get("SLACK_SIGNING_SECRET")
        if not signing_secret:
            raise HTTPException(status_code=503, detail="Slack is not configured")
        body = await request.body()
        if not verify_signature(
            body,
            request.headers.get("X-Slack-Request-Timestamp", ""),
            request.headers.get("X-Slack-Signature", ""),
            signing_secret,
        ):
            files.record_event(
                "connection.event_rejected",
                "unknown",
                "connector",
                "slack",
                details={"reason": "signature"},
                outcome="failed",
            )
            raise HTTPException(status_code=403, detail="Invalid Slack signature")
        try:
            payload = await request.json()
        except ValueError as error:
            raise HTTPException(status_code=400, detail="Invalid Slack event") from error
        if not isinstance(payload, dict):
            raise HTTPException(status_code=400, detail="Invalid Slack event")
        if payload.get("type") == "url_verification":
            return JSONResponse({"challenge": payload.get("challenge", "")})
        delivery = ingest_slack_event(payload)
        if delivery:
            background.add_task(analyze_slack, *delivery)
        return JSONResponse({"ok": True})

    return app


app = create_app()
